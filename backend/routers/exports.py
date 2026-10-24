import os
import tempfile
from contextlib import suppress
from datetime import datetime
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from auth import get_current_user
from database import get_db
from models import Application, Customer, OperationLog
from utils.masking import mask_id_number, mask_phone

router = APIRouter(prefix="/api/exports", tags=["数据导出"])


def _log_export(db, current_user: dict, action: str, detail: str = "") -> None:
    """Write an export audit log.

    Exports are the primary channel for full data leaving the system, so they must be
    traceable: who, when, what was exported, and whether masking was applied.
    Failures do not affect the export itself (auditing is an auxiliary capability).
    """
    try:
        db.add(OperationLog(
            user_id=current_user.get("user_id") or current_user.get("id"),
            username=current_user.get("username", ""),
            action=action,
            resource_type="export",
            new_value={"detail": detail} if detail else None,
        ))
        db.commit()
    except Exception:
        db.rollback()

EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Chunk size for streaming the temp file
STREAM_CHUNK_SIZE = 64 * 1024
# Max rows per export; anything beyond is truncated and noted in the footer row
MAX_EXPORT_ROWS = 50000
# Number of records prefetched per step by the database streaming cursor
YIELD_PER = 500

# Column widths are fixed values: write_only mode cannot iterate ws.columns after writing to compute them
CUSTOMER_COLUMN_WIDTHS = [14, 22, 14, 12, 16, 10, 28, 16, 10, 20, 16, 18]
APPLICATION_COLUMN_WIDTHS = [22, 14, 22, 16, 18, 18]


def _set_column_widths(ws, widths):
    """Column widths must be set before writing data in write_only mode."""
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width


def _safe_remove(path):
    with suppress(OSError):
        os.remove(path)


def _iter_file(path, chunk_size=STREAM_CHUNK_SIZE):
    """Read the temp file in chunks; the file is removed in finally whether the stream ends normally or the client disconnects."""
    try:
        with open(path, "rb") as f:
            while True:
                chunk = f.read(chunk_size)
                if not chunk:
                    break
                yield chunk
    finally:
        _safe_remove(path)


def _stream_write_only_workbook(title, build_rows, disposition):
    """Write the write_only workbook to a temp file, then stream it back in chunks.

    build_rows(ws) is responsible for appending rows to the write_only worksheet.
    The workbook is fully written to disk before the response is returned, so the
    streaming phase no longer touches the database session. The temp file is cleaned
    up with double protection: the generator's finally + a BackgroundTask fallback
    (in case the response object is never consumed).
    """
    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", prefix="export_")
    os.close(fd)
    try:
        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet(title=title)
        build_rows(ws)
        wb.save(tmp_path)
    except Exception:
        _safe_remove(tmp_path)
        raise

    return StreamingResponse(
        _iter_file(tmp_path),
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": disposition},
        background=BackgroundTask(_safe_remove, tmp_path),
    )


@router.get("/template")
async def download_excel_template(request: Request):
    """Download the Excel bulk import template"""
    user = await get_current_user(request)
    if user.get("role") not in ("admin", "salesman"):
        raise HTTPException(status_code=403, detail="仅业务员和管理员可下载模板")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "批量导入模板"

    headers = [
        "客户姓名", "身份证号", "手机号", "学历",
        "现职称", "工作单位", "岗位",
    ]
    ws.append(headers)

    # Style the header row
    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Set column widths
    widths = [14, 22, 14, 16, 14, 22, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Add one sample row (comment row)
    ws.append(["张三", "110101199001011234", "13800138000", "本科", "中级工程师", "某某科技有限公司", "技术主管"])

    from urllib.parse import quote
    filename = quote("Excel批量导入模板.xlsx", safe='')
    ascii_name = "batch_import_template.xlsx"
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.read(),
        media_type=EXCEL_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{filename}',
        },
    )


@router.post("/customers")
async def export_customers(
    status: str = Query(None),
    keyword: str = Query(None),
    mask: bool = Query(False, description="是否脱敏身份证号与手机号（对外发送建议开启）"),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    if current_user.get("role") not in ("admin", "salesman"):
        raise HTTPException(status_code=403, detail="仅业务员和管理员可导出")

    query = db.query(Customer, Application).join(
        Application, Customer.id == Application.customer_id
    )

    if current_user.get("role") == "salesman":
        query = query.filter(Customer.assigned_salesman_id == current_user.get("id"))

    if status:
        query = query.filter(Application.status == status)

    if keyword:
        query = query.filter(
            or_(
                Customer.name.like(f"%{keyword}%"),
                Customer.id_number.like(f"%{keyword}%"),
                Customer.phone.like(f"%{keyword}%"),
                Customer.work_unit.like(f"%{keyword}%"),
            )
        )

    query = query.order_by(Application.id.desc())

    headers = [
        "客户姓名", "身份证号", "手机号", "学历", "现职称", "获聘年份",
        "工作单位", "岗位", "专业年限", "申报批次", "当前状态", "创建时间",
    ]

    def build_rows(ws):
        _set_column_widths(ws, CUSTOMER_COLUMN_WIDTHS)
        ws.append(headers)

        written = 0
        truncated = False
        # Fetch 1 extra row to detect truncation; yield_per uses a streaming cursor instead of loading everything into memory
        rows = query.limit(MAX_EXPORT_ROWS + 1).yield_per(YIELD_PER)
        for customer, app in rows:
            if written >= MAX_EXPORT_ROWS:
                truncated = True
                continue
            ws.append([
                customer.name,
                mask_id_number(customer.id_number) if mask else customer.id_number + "\t",
                mask_phone(customer.phone) if mask else (customer.phone or ""),
                customer.education or "",
                customer.current_title or "",
                customer.current_title_year or "",
                customer.work_unit or "",
                customer.position or "",
                customer.professional_years or "",
                app.batch_number if app else "",
                app.status if app else "",
                customer.created_at.strftime("%Y-%m-%d %H:%M") if customer.created_at else "",
            ])
            written += 1

        if truncated:
            ws.append([f"（数据量超限，仅导出前 {MAX_EXPORT_ROWS} 条）"])

    from urllib.parse import quote
    filename = quote(f"客户列表_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    # Export audit: full customer data (including ID numbers) is exported and must be traceable
    _log_export(
        db, current_user, "导出客户列表",
        detail=f"脱敏={'是' if mask else '否'}, 状态筛选={status or '全部'}, 关键词={keyword or '无'}",
    )

    return _stream_write_only_workbook(
        "客户列表",
        build_rows,
        f'attachment; filename="{filename}"',
    )


@router.post("/applications")
async def export_applications(
    status: str = Query(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    if current_user.get("role") not in ("admin", "salesman"):
        raise HTTPException(status_code=403, detail="仅业务员和管理员可导出")

    query = db.query(Application, Customer).join(
        Customer, Application.customer_id == Customer.id
    )

    if current_user.get("role") == "salesman":
        query = query.filter(Customer.assigned_salesman_id == current_user.get("id"))

    if status:
        query = query.filter(Application.status == status)

    query = query.order_by(Application.id.desc())

    headers = [
        "批次号", "客户姓名", "客户 ID 号", "申报状态", "创建时间", "更新时间",
    ]

    def build_rows(ws):
        _set_column_widths(ws, APPLICATION_COLUMN_WIDTHS)
        ws.append(headers)

        written = 0
        truncated = False
        rows = query.limit(MAX_EXPORT_ROWS + 1).yield_per(YIELD_PER)
        for app, customer in rows:
            if written >= MAX_EXPORT_ROWS:
                truncated = True
                continue
            ws.append([
                app.batch_number,
                customer.name,
                customer.id_number + "\t",
                app.status,
                app.created_at.strftime("%Y-%m-%d %H:%M") if app.created_at else "",
                app.updated_at.strftime("%Y-%m-%d %H:%M") if app.updated_at else "",
            ])
            written += 1

        if truncated:
            ws.append([f"（数据量超限，仅导出前 {MAX_EXPORT_ROWS} 条）"])

    from urllib.parse import quote
    filename = quote(f"申报批次_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    _log_export(db, current_user, "导出申报批次", detail=f"状态筛选={status or '全部'}")

    return _stream_write_only_workbook(
        "申报批次",
        build_rows,
        f'attachment; filename="{filename}"',
    )
