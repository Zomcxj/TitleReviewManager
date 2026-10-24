import json
import os
import tempfile
from contextlib import suppress
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import desc
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from auth import get_current_user
from database import get_db
from models import OperationLog
from schemas import OperationLogListResponse

router = APIRouter(prefix="/api/audit", tags=["审计日志"])

# Max rows per export; anything beyond is truncated and noted in the footer row
EXPORT_MAX_ROWS = 10000
# Max JSON text length per cell, to avoid overly long Excel cells
EXPORT_CELL_MAX_LEN = 1000

EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Chunk size for streaming the temp file
STREAM_CHUNK_SIZE = 64 * 1024
# Number of records prefetched per step by the database streaming cursor
YIELD_PER = 500
# Fixed column widths: write_only mode cannot iterate ws.columns after writing to compute them
EXPORT_COLUMN_WIDTHS = [8, 14, 28, 14, 10, 40, 40, 16, 30, 20]
# Audit export headers (write_only mode requires per-cell styling via WriteOnlyCell)
EXPORT_HEADERS = [
    "ID", "操作人", "操作类型", "资源类型", "资源ID",
    "变更前", "变更后", "IP地址", "User-Agent", "操作时间",
]


def _safe_remove(path):
    with suppress(OSError):
        os.remove(path)


def _iter_file(path, chunk_size=STREAM_CHUNK_SIZE):
    """Read the temp file in chunks; the file is removed in finally whether the stream ends normally or the client disconnects."""
    try:
        with open(path, "rb") as f:
            while True:
                chunk = f.read(chunk_size)
                if not chunk:
                    break
                yield chunk
    finally:
        _safe_remove(path)


def _build_log_query(
    db: Session,
    action: str | None,
    resource_type: str | None,
    username: str | None,
    start_date: datetime | None,
    end_date: datetime | None,
):
    """Build the audit log query from filter criteria (shared by list and export endpoints for consistent behavior)."""
    query = db.query(OperationLog)

    if action:
        query = query.filter(OperationLog.action.like(f"%{action}%"))
    if resource_type:
        query = query.filter(OperationLog.resource_type == resource_type)
    if username:
        query = query.filter(OperationLog.username.like(f"%{username}%"))
    if start_date:
        query = query.filter(OperationLog.created_at >= start_date)
    if end_date:
        query = query.filter(OperationLog.created_at <= end_date)

    return query


@router.get("/export")
async def export_operation_logs(
    action: str | None = None,
    resource_type: str | None = None,
    username: str | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Export audit logs to Excel (admin only, up to 10000 rows).

    Uses openpyxl write_only mode + SQLAlchemy yield_per streaming cursor
    to avoid loading all logs into memory at once; the workbook is written to a
    temp file and streamed back in chunks.
    """
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="只有管理员可以导出审计日志")

    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    query = _build_log_query(db, action, resource_type, username, start_date, end_date)
    # Fetch 1 extra row to detect truncation, saving a count query
    logs = (
        query.order_by(desc(OperationLog.created_at))
        .limit(EXPORT_MAX_ROWS + 1)
        .yield_per(YIELD_PER)
    )

    def _serialize(value):
        if value is None:
            return ""
        try:
            text = json.dumps(value, ensure_ascii=False, default=str)
        except Exception:
            text = str(value)
        return text[:EXPORT_CELL_MAX_LEN]

    def _build(ws):
        # write_only mode: column widths must be set before writing data, and ws[1] cannot be used to iterate cells
        for i, width in enumerate(EXPORT_COLUMN_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_cells = []
        for title in EXPORT_HEADERS:
            cell = WriteOnlyCell(ws, value=title)
            cell.font = header_font
            cell.fill = header_fill
            header_cells.append(cell)
        ws.append(header_cells)

        written = 0
        truncated = False
        for log in logs:
            if written >= EXPORT_MAX_ROWS:
                truncated = True
                continue
            ws.append([
                log.id,
                log.username or "",
                log.action or "",
                log.resource_type or "",
                log.resource_id if log.resource_id is not None else "",
                _serialize(log.old_value),
                _serialize(log.new_value),
                log.ip_address or "",
                log.user_agent or "",
                log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else "",
            ])
            written += 1

        if truncated:
            ws.append([f"（已截断，仅导出前 {EXPORT_MAX_ROWS} 条）"])

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", prefix="audit_export_")
    os.close(fd)
    try:
        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet(title="操作审计日志")
        _build(ws)
        wb.save(tmp_path)
    except Exception:
        _safe_remove(tmp_path)
        raise

    # The audit log export itself must be logged (audit data leaving the system must also be traceable)
    try:
        db.add(OperationLog(
            user_id=current_user.get("user_id") or current_user.get("id"),
            username=current_user.get("username", ""),
            action="导出审计日志",
            resource_type="export",
            new_value={"detail": f"筛选: action={action or '全部'}, resource_type={resource_type or '全部'}"},
        ))
        db.commit()
    except Exception:
        db.rollback()

    filename = quote(f"操作审计日志_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx", safe="")
    return StreamingResponse(
        _iter_file(tmp_path),
        media_type=EXCEL_MEDIA_TYPE,
        headers={
            "Content-Disposition": f"attachment; filename=\"audit_logs.xlsx\"; filename*=UTF-8''{filename}",
        },
        background=BackgroundTask(_safe_remove, tmp_path),
    )


@router.get("", response_model=OperationLogListResponse)
@router.get("/logs", response_model=OperationLogListResponse)
async def get_operation_logs(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    action: str | None = None,
    resource_type: str | None = None,
    username: str | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="只有管理员可以查看审计日志")
    
    query = _build_log_query(db, action, resource_type, username, start_date, end_date)
    
    total = query.count()
    logs = (
        query.order_by(desc(OperationLog.created_at))
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    
    return {
        "items": logs,
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.get("/logs/resource/{resource_type}/{resource_id}")
async def get_resource_logs(
    resource_type: str,
    resource_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="只有管理员可以查看审计日志")
    
    logs = (
        db.query(OperationLog)
        .filter(
            OperationLog.resource_type == resource_type,
            OperationLog.resource_id == resource_id,
        )
        .order_by(desc(OperationLog.created_at))
        .limit(50)
        .all()
    )
    
    return {"items": logs, "total": len(logs)}


@router.get("/verify-chain")
async def verify_audit_chain(
    limit: int = Query(None, ge=1, le=100000, description="只校验最近 N 条，不传则全量"),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Verify audit log hash chain integrity (admin only).

    Detects whether logs were tampered with or deleted: any modification breaks
    the rest of the chain. Returns broken_at listing the ids of logs that failed
    verification, for easy locating.
    """
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="只有管理员可以校验审计日志")

    from utils.audit_chain import verify_chain
    result = verify_chain(db, limit=limit)
    if not result["valid"]:
        # Tampering is a serious event; write an audit record (this record is itself on the chain, so attackers cannot erase history)
        db.add(OperationLog(
            user_id=current_user.get("user_id") or current_user.get("id"),
            username=current_user.get("username", ""),
            action="审计日志链校验失败",
            resource_type="audit",
            new_value={"detail": result.get("reason"), "broken_at": result.get("broken_at")},
        ))
        db.commit()
    return result


@router.post("/backfill-chain")
async def backfill_audit_chain(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Backfill hash chains for historical logs (admin only; call once after the feature ships)."""
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="只有管理员可以操作")

    from utils.audit_chain import backfill_chain
    total = 0
    while True:
        n = backfill_chain(db, batch_size=500)
        total += n
        if n == 0:
            break
    return {"message": f"已为 {total} 条历史日志补算哈希链", "count": total}
